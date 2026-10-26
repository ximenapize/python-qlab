"""Auditoría humana del LLM: sortea 20 extracciones (10 documentos y 10 audios) y arma la plantilla Excel.

    .venv/bin/python auditoria.py

El sorteo usa solo los ids y una semilla fija (no mira las extracciones). La plantilla va a
datos/raw/auditoria_llm.xlsx (fuera de Git mientras se llena); 02_evaluacion.ipynb la lee y
publica el resultado en datos/auditoria_llm.csv.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CARPETA = Path(__file__).resolve().parent
PLANTILLA = CARPETA / "datos" / "raw" / "auditoria_llm.xlsx"
SEMILLA = 2026
N_POR_MODALIDAD = 10
CAMPOS = {
    "documento": ["tema", "resumen", "emisor", "fecha", "monto"],
    "audio": ["tema", "resumen", "citas_clave", "numero_hablantes"],
}

CRITERIOS = [
    ("Regla general",
     "Se juzga contra el TEXTO que recibió el LLM (OCR o transcripción, columna texto_fuente), no contra la imagen o el "
     "audio original: se evalúa al LLM, no al OCR/ASR. Si el texto trae un error de OCR y el LLM lo copia tal cual, el "
     "campo es correcto."),
    ("tema", "Correcto si nombra el asunto principal del texto en pocas palabras. Incorrecto si el asunto es otro, es "
             "demasiado genérico para distinguir el texto (p. ej. «documento histórico») o agrega algo que no está."),
    ("resumen", "Correcto si es fiel al texto y no agrega hechos, nombres, lugares, fechas ni cifras que no aparezcan. "
                "Un solo dato inventado o una afirmación contraria al texto lo hace incorrecto. No se exige que sea completo."),
    ("emisor", "Correcto si el valor aparece en el texto y es quien emite o firma (institución, periódico, autor), o si "
               "es null y el texto no dice quién emite. Incorrecto si es inventado, si es otra entidad mencionada o si "
               "es null cuando el emisor aparece claramente."),
    ("fecha", "Correcto si es una fecha escrita en el texto (puede venir con errores de OCR copiados) o null si no hay "
              "ninguna. Incorrecto si es inventada, deducida o si es null habiendo una fecha clara."),
    ("monto", "Correcto si es un importe de dinero que aparece en el texto, o null si no hay ninguno. Incorrecto si es "
              "un número que no es dinero (año, número de página, artículo), si es inventado o si falta habiendo uno claro."),
    ("citas_clave", "Correcto si cada cita aparece literalmente en el texto (se toleran mayúsculas, puntuación y las "
                    "marcas [HABLANTE_N]) y es relevante para el tema. Incorrecto si alguna cita no está en el texto o "
                    "está parafraseada. null es correcto solo si el texto es demasiado corto para tener citas."),
    ("numero_hablantes", "Correcto si coincide con el número de marcas [HABLANTE_N] distintas del texto, o si es null "
                         "cuando el texto no tiene marcas. Incorrecto en cualquier otro caso."),
    ("Extracción completamente correcta",
     "Sí, solo si TODOS los campos de la fila están marcados como «correcto». Se calcula sola en la columna "
     "completamente_correcta. El resultado final es «N de 20 extracciones completamente correctas»."),
]


def sortear(extraccion: pd.DataFrame) -> pd.DataFrame:
    return pd.concat([
        extraccion[extraccion["modalidad"] == m].sample(N_POR_MODALIDAD, random_state=SEMILLA).sort_values("id")
        for m in ["documento", "audio"]
    ])


def textos_fuente() -> dict[str, str]:
    docs = pd.read_csv(CARPETA / "datos" / "documentos.csv")
    clips = pd.read_csv(CARPETA / "datos" / "transcripciones.csv")
    return {**dict(zip(docs["id"], docs["texto_ocr"].fillna(""))), **dict(zip(clips["id"], clips["texto_asr"].fillna("")))}


def archivos() -> dict[str, str]:
    f = pd.read_csv(CARPETA / "datos" / "fuentes.csv")
    return dict(zip(f["id"], f["archivo"]))


def _valor(v) -> str:
    if v is None:
        return "null"
    if isinstance(v, list):
        return "\n".join(f"• {x}" for x in v)
    return str(v)


def hoja(wb: Workbook, nombre: str, muestra: pd.DataFrame, campos: list[str], texto: dict, archivo: dict):
    ws = wb.create_sheet(nombre)
    cols = ["id", "archivo", "texto_fuente", "json_llm"]
    for c in campos:
        cols += [f"{c}: valor", f"{c}: ¿correcto?"]
    cols += ["completamente_correcta", "notas"]
    ws.append(cols)
    dv = DataValidation(type="list", formula1='"correcto,incorrecto"', allow_blank=True)
    ws.add_data_validation(dv)
    for f in muestra.itertuples():
        c = json.loads(f.campos_json)
        fila = [f.id, archivo.get(f.id, ""), texto.get(f.id, ""), json.dumps(c, ensure_ascii=False, indent=1)]
        for k in campos:
            fila += [_valor(c.get(k)), ""]
        ws.append(fila + ["", ""])
        r = ws.max_row
        ws.cell(r, 2).hyperlink = (CARPETA / archivo.get(f.id, "")).as_uri()
        ws.cell(r, 2).font = Font(color="2A78D6", underline="single")
        celdas = [f"{get_column_letter(5 + 2 * i + 1)}{r}" for i in range(len(campos))]
        for cel in celdas:
            dv.add(cel)
        todas = ",".join(f'{cel}="correcto"' for cel in celdas)
        vacias = "+".join(f'({cel}="")' for cel in celdas)
        ws.cell(r, len(cols) - 1).value = f'=IF({vacias}>0,"",IF(AND({todas}),"sí","no"))'
    # formato
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0CA30C" if "¿correcto?" in str(cel.value) else "2A78D6")
    anchos = {"id": 10, "archivo": 22, "texto_fuente": 60, "json_llm": 45, "completamente_correcta": 14, "notas": 30}
    for i, nombre_col in enumerate(cols, start=1):
        ancho = anchos.get(nombre_col, 14 if "¿correcto?" in nombre_col else 30)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for cel in fila:
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila[0].row].height = 220
    rango = f"A2:{get_column_letter(len(cols))}{ws.max_row}"
    col_ok = get_column_letter(len(cols) - 1)
    ws.conditional_formatting.add(rango, FormulaRule(formula=[f'${col_ok}2="sí"'], fill=PatternFill("solid", fgColor="E3F4E3")))
    ws.conditional_formatting.add(rango, FormulaRule(formula=[f'${col_ok}2="no"'], fill=PatternFill("solid", fgColor="FBE3E3")))
    ws.freeze_panes = "C2"


def crear_plantilla() -> pd.DataFrame:
    extraccion = pd.read_csv(CARPETA / "datos" / "extraccion_llm.csv")
    muestra = sortear(extraccion)
    texto, archivo = textos_fuente(), archivos()
    wb = Workbook()
    ins = wb.active
    ins.title = "criterios"
    ins.append(["Criterios de corrección para la auditoría humana del LLM (20 extracciones)"])
    ins["A1"].font = Font(bold=True, size=13)
    ins.append([f"Muestra sorteada con auditoria.py (semilla {SEMILLA}, 10 documentos y 10 audios) sin mirar las extracciones."])
    ins.append(["En cada hoja, elige «correcto» o «incorrecto» en las columnas verdes. La columna completamente_correcta se llena sola."])
    ins.append([])
    ins.append(["Campo", "Criterio"])
    for c in ins[5]:
        c.font = Font(bold=True)
    for campo, criterio in CRITERIOS:
        ins.append([campo, criterio])
    ins.column_dimensions["A"].width = 28
    ins.column_dimensions["B"].width = 120
    for fila in ins.iter_rows(min_row=6):
        for cel in fila:
            cel.alignment = Alignment(wrap_text=True, vertical="top")
    for m, nombre in [("documento", "documentos"), ("audio", "audios")]:
        hoja(wb, nombre, muestra[muestra["modalidad"] == m], CAMPOS[m], texto, archivo)
    PLANTILLA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PLANTILLA)
    return muestra


if __name__ == "__main__":
    m = crear_plantilla()
    print(f"Plantilla: {PLANTILLA.relative_to(CARPETA)} · {len(m)} extracciones: {', '.join(m['id'])}")
